"""
excel_to_json.py - convert a questions Excel/CSV file into the app's JSON format.

Usually you do NOT need this: the website itself can read an .xlsx file on the
"Questions" page. Use this script when you want to keep the JSON inside the
repository (data/*.json) so every visitor gets the questions automatically.

Usage:
    python tools/excel_to_json.py questions.xlsx
    python tools/excel_to_json.py questions.xlsx --out data/my-subject.json
    python tools/excel_to_json.py questions.xlsx --out data/mixed.json --append
"""

import argparse
import csv
import json
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

LETTERS = ['A', 'B', 'C', 'D', 'E', 'F']

FIELD_ALIASES = {
    'id': ['id', 'qid', 'question_id', 'sr', 's.no', 'sno', 'no', 'q.no', 'number'],
    'subject': ['subject', 'sub', 'section', 'paper', 'category'],
    'topic': ['topic', 'chapter', 'unit', 'subtopic'],
    'difficulty': ['difficulty', 'level'],
    'q_hi': ['q_hi', 'question_hi', 'question_hindi', 'hindi', 'que_hi'],
    'q_en': ['q_en', 'question_en', 'question_english', 'english', 'que_en'],
    'answer': ['answer', 'ans', 'correct', 'correct_option', 'correct_answer', 'key', 'answer_key'],
    'expl_hi': ['expl_hi', 'explanation_hi', 'explain_hi', 'sol_hi'],
    'expl_en': ['expl_en', 'explanation_en', 'explain_en', 'sol_en', 'explanation', 'solution'],
    'tags': ['tags', 'keywords'],
}


def norm_key(k):
    return str(k or '').strip().lower().replace(' ', '_')


def pick(look, names):
    for n in names:
        v = look.get(n)
        if v is not None and str(v).strip() != '':
            return str(v).strip()
    return ''


def answer_index(ans, count):
    if not ans:
        return -1
    s = str(ans).strip()
    up = s.upper()
    if len(up) == 1 and up in LETTERS:
        i = LETTERS.index(up)
        return i if i < count else -1
    if s.isdigit():
        n = int(s)
        if 1 <= n <= count:
            return n - 1
        if n == 0:
            return 0
    return -1


def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.csv', '.txt'):
        with open(path, 'r', encoding='utf-8-sig', newline='') as f:
            return [{norm_key(k): v for k, v in row.items()} for row in csv.DictReader(f)]
    if load_workbook is None:
        sys.exit('openpyxl is required for Excel files:  pip install openpyxl')
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        if ws.title.lower().startswith('how-to'):
            continue
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(v is None or str(v).strip() == '' for v in row):
                continue
            rows.append({norm_key(header[i]): row[i] for i in range(min(len(header), len(row)))})
    return rows



def convert(rows, subject_default='General'):
    out, problems = [], []
    for n, look in enumerate(rows, start=2):
        opts_hi, opts_en = [], []
        for i in range(1, 7):
            h = pick(look, ['opt%d_hi' % i, 'option%d_hi' % i, 'opt_%d_hi' % i, 'choice%d_hi' % i])
            e = pick(look, ['opt%d_en' % i, 'option%d_en' % i, 'opt_%d_en' % i, 'choice%d_en' % i])
            p = pick(look, ['opt%d' % i, 'option%d' % i, 'opt_%d' % i, 'choice%d' % i])
            opts_hi.append(h or p)
            opts_en.append(e)
        while opts_hi and not opts_hi[-1]:
            opts_hi.pop()
        while opts_en and not opts_en[-1]:
            opts_en.pop()
        if not opts_en:
            opts_en = list(opts_hi)
        if not opts_hi:
            opts_hi = list(opts_en)

        q_hi = pick(look, FIELD_ALIASES['q_hi'])
        q_en = pick(look, FIELD_ALIASES['q_en'])
        count = max(len(opts_hi), len(opts_en))
        ai = answer_index(pick(look, FIELD_ALIASES['answer']), count)

        if not (q_hi or q_en):
            problems.append('row %d: no question text' % n)
            continue
        if count < 2:
            problems.append('row %d: fewer than 2 options' % n)
            continue
        if ai < 0:
            problems.append('row %d: invalid or missing answer' % n)
            continue

        out.append({
            'id': pick(look, FIELD_ALIASES['id']) or ('excel-%d' % n),
            'subject': pick(look, FIELD_ALIASES['subject']) or subject_default,
            'topic': pick(look, FIELD_ALIASES['topic']) or 'General',
            'difficulty': (pick(look, FIELD_ALIASES['difficulty']) or 'medium').lower(),
            'question': {'hi': q_hi, 'en': q_en},
            'options': {'hi': opts_hi, 'en': opts_en},
            'answer': LETTERS[ai],
            'explanation': {
                'hi': pick(look, FIELD_ALIASES['expl_hi']),
                'en': pick(look, FIELD_ALIASES['expl_en']),
            },
            'tags': [t.strip() for t in pick(look, FIELD_ALIASES['tags']).split(',') if t.strip()],
        })
    return out, problems


def main():
    ap = argparse.ArgumentParser(description='Convert an Excel/CSV question file to app JSON.')
    ap.add_argument('input', help='path to .xlsx / .csv file')
    ap.add_argument('--out', help='output json path (default: data/<same-name>.json)')
    ap.add_argument('--append', action='store_true', help='merge into an existing json file')
    args = ap.parse_args()

    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(here)
    rows = read_rows(args.input)
    questions, problems = convert(rows)

    out = args.out or os.path.join(root, 'data', os.path.splitext(os.path.basename(args.input))[0] + '.json')
    if args.append and os.path.exists(out):
        with open(out, 'r', encoding='utf-8') as f:
            existing = json.load(f)
        merged = {}
        for q in list(existing) + questions:
            merged[q['question']['hi'] or q['question']['en']] = q
        questions = list(merged.values())

    with open(out, 'w', encoding='utf-8') as f:
        json.dump(questions, f, ensure_ascii=False, indent=2)

    print('read %d row(s), wrote %d question(s) to %s' % (len(rows), len(questions), out))
    if problems:
        print('\nskipped rows:')
        for p in problems[:20]:
            print('  -', p)
        if len(problems) > 20:
            print('  ... and %d more' % (len(problems) - 20))
    print('\nRemember: add the file name to data/index.json so the website loads it.')


if __name__ == '__main__':
    main()

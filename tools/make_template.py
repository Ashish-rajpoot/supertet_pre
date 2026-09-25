"""
make_template.py - create templates/questions-template.xlsx

One row = one question. Only `answer` is compulsory alongside the question text.
Columns starting with q_ / expl_ are widened so Hindi text is readable.

Run:  python tools/make_template.py
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    'id', 'subject', 'topic', 'difficulty',
    'q_hi', 'q_en',
    'opt1_hi', 'opt2_hi', 'opt3_hi', 'opt4_hi',
    'opt1_en', 'opt2_en', 'opt3_en', 'opt4_en',
    'answer', 'expl_hi', 'expl_en', 'tags',
]

ROWS = [
    {
        'id': 'gk-101', 'subject': 'GK & GS', 'topic': 'Important Days', 'difficulty': 'easy',
        'q_hi': 'राष्ट्रीय युवा दिवस कब मनाया जाता है?',
        'q_en': 'When is National Youth Day celebrated?',
        'opt1_hi': '10 जनवरी', 'opt2_hi': '12 जनवरी', 'opt3_hi': '15 जनवरी', 'opt4_hi': '24 जनवरी',
        'opt1_en': '10 January', 'opt2_en': '12 January', 'opt3_en': '15 January', 'opt4_en': '24 January',
        'answer': 'B',
        'expl_hi': '12 जनवरी को स्वामी विवेकानंद का जन्मदिन है।',
        'expl_en': '12 January is Swami Vivekananda birth anniversary.',
        'tags': 'days,national',
    },
    {
        'id': 'sci-101', 'subject': 'Science', 'topic': 'Human Body', 'difficulty': 'medium',
        'q_hi': 'मानव शरीर में कुल कितनी हड्डियाँ होती हैं?',
        'q_en': 'How many bones are there in the human body?',
        'opt1_hi': '206', 'opt2_hi': '208', 'opt3_hi': '204', 'opt4_hi': '210',
        'opt1_en': '206', 'opt2_en': '208', 'opt3_en': '204', 'opt4_en': '210',
        'answer': 'A', 'expl_hi': '', 'expl_en': '', 'tags': 'human body',
    },
    {
        'id': 'hin-101', 'subject': 'Hindi', 'topic': 'व्याकरण', 'difficulty': 'easy',
        'q_hi': "'विद्यालय' शब्द में कौन-सी संधि है?",
        'q_en': '',
        'opt1_hi': 'दीर्घ संधि', 'opt2_hi': 'गुण संधि', 'opt3_hi': 'वृद्धि संधि', 'opt4_hi': 'यण संधि',
        'opt1_en': '', 'opt2_en': '', 'opt3_en': '', 'opt4_en': '',
        'answer': 'A', 'expl_hi': 'विद्या + आलय = विद्यालय', 'expl_en': '', 'tags': 'sandhi',
    },
]

HELP = [
    ('id', 'Any unique text. Leave blank and the app will generate one.'),
    ('subject', 'Used to build the test menu, e.g. GK & GS / Science / Hindi.'),
    ('topic', 'Smaller heading inside the subject, e.g. Important Days.'),
    ('difficulty', 'easy / medium / hard (optional, defaults to medium).'),
    ('q_hi / q_en', 'Question text in Hindi and English. Fill at least one.'),
    ('opt1_hi..opt4_hi', 'Four options in Hindi.'),
    ('opt1_en..opt4_en', 'Four options in English. If left blank, Hindi options are reused.'),
    ('answer', 'A, B, C or D. You can also write 1, 2, 3, 4.'),
    ('expl_hi / expl_en', 'Optional explanation shown after the answer.'),
    ('tags', 'Optional comma separated keywords.'),
]


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    out_dir = os.path.join(os.path.dirname(here), 'templates')
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, 'questions-template.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.title = 'questions'

    head_fill = PatternFill('solid', fgColor='2F6DF6')
    head_font = Font(color='FFFFFF', bold=True)
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(c)].width = (
            44 if name.startswith(('q_', 'expl_')) else (20 if name.startswith('opt') else 16)
        )
    ws.freeze_panes = 'A2'

    for r, row in enumerate(ROWS, start=2):
        for c, name in enumerate(COLUMNS, start=1):
            ws.cell(row=r, column=c, value=row.get(name, ''))

    help_ws = wb.create_sheet('how-to-fill')
    help_ws['A1'] = 'Column'
    help_ws['B1'] = 'Meaning'
    help_ws['A1'].font = Font(bold=True)
    help_ws['B1'].font = Font(bold=True)
    help_ws.column_dimensions['A'].width = 24
    help_ws.column_dimensions['B'].width = 80
    for i, (col, note) in enumerate(HELP, start=2):
        help_ws.cell(row=i, column=1, value=col)
        help_ws.cell(row=i, column=2, value=note)

    wb.save(path)
    print('wrote', path)


if __name__ == '__main__':
    main()
